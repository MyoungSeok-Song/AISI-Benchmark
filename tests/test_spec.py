"""명세 생성 모듈(kyab_runner/spec) 검사: 자료 형태, 원본 xlsx와의 대응, 추출기의 재현성.

원본 xlsx는 저장소 밖(project proposal/)에 있어 없을 수 있다. 있으면 추출기를 다시 돌려 생성 모듈과
바이트 단위로 같은지 본다(원본이 바뀌지 않았는데 모듈이 손으로 고쳐졌거나, 추출기가 달라졌으면 어긋난다).
실행 (runner/ 폴더에서):  .venv/bin/python -m unittest discover -s tests -v
"""
import importlib.util
import io
import os
import tempfile
import unittest
from unittest import mock
from pathlib import Path

from test_runner import CODEBOOK, RUNNER_DIR, TAXONOMY

import _xlsx_common                                                        # noqa: E402
import extract_codebook                                                    # noqa: E402
import extract_taxonomy                                                    # noqa: E402
from kyab_runner.spec import codebook_data, taxonomy_data                  # noqa: E402

# 파일명은 NFC로 비교한다(macOS에서 온 파일은 NFD). 없으면 None → 재현성 검사는 건너뜀
CODEBOOK_XLSX = _xlsx_common.locate_xlsx(_xlsx_common.DEFAULT_SRC_DIR, codebook_data.SOURCE_FILE)
TAXONOMY_XLSX = _xlsx_common.locate_xlsx(_xlsx_common.DEFAULT_SRC_DIR, taxonomy_data.TAXONOMY["source_file"])


def load_module(path):
    spec = importlib.util.spec_from_file_location(path.stem, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


class SpecDataTest(unittest.TestCase):
    def test_codebook_module_shape(self):
        """7개 표, 시트가 선언한 필드 수와 같고, 로더가 만든 Codebook의 열과 순서가 같다(overlay가 더한 열 제외)."""
        tables = codebook_data.CODEBOOK["tables"]
        self.assertEqual(list(tables), ["01_items", "02_item_tags", "03_prompts", "04_runs",
                                        "05_responses", "06_judgments", "07_results"])
        for name, table in tables.items():
            self.assertEqual(table["declared_field_count"], len(table["fields"]), name)
            names = [f["name"] for f in table["fields"]]
            self.assertEqual(names, [c for c in CODEBOOK.columns(name) if c in names], name)
            self.assertEqual(set(table["fields"][0]), {*(k for _, k in extract_codebook.HEADER_LABELS), "enum_kind", "enum", "regex"})
        self.assertEqual(codebook_data.CODEBOOK_VERSION, extract_codebook.CODEBOOK_VERSION)
        self.assertEqual(codebook_data.CODEBOOK["codebook_version"], codebook_data.CODEBOOK_VERSION)
        self.assertRegex(codebook_data.SOURCE_SHA256, r"^[0-9a-f]{64}$")

    def test_taxonomy_module_shape(self):
        """상위 10 · 하위 39, 대응표 49행(상위 10 + 하위 39), 로더 결과와 일치."""
        majors = taxonomy_data.TAXONOMY["majors"]
        self.assertEqual(len(majors), 10)
        self.assertEqual(sum(len(m["subs"]) for m in majors), 39)
        self.assertEqual(len(taxonomy_data.CROSSWALK), 49)
        self.assertEqual(list(taxonomy_data.CROSSWALK[0]), ["legacy_code", "legacy_scheme", "new_code", "level",
                                                            "relation", "new_name", "source_cell"])
        self.assertEqual(TAXONOMY.major_codes, tuple(m["code"] for m in majors))
        self.assertEqual((TAXONOMY.legacy_risk_codes, TAXONOMY.legacy_m_codes),
                         (("R1", "R2", "R3", "R4", "R5"), ("M01", "M02", "M03", "M04", "M05")))

    def test_render_module_round_trips_and_is_deterministic(self):
        """생성 모듈 본문은 실행하면 같은 자료가 되고(한글·정수·None 포함), 같은 입력이면 같은 바이트다."""
        data = {"가": [1, 2, {"나": None, "다": "긴 문장 " * 30}], "enum": [1, 2, 3], "x": "y"}
        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "원본.xlsx"
            source.write_bytes(b"xlsx")
            text = _xlsx_common.render_module("설명", source, "python tools/x.py", [("DATA", data), ("N", 3)])
            self.assertEqual(text, _xlsx_common.render_module("설명", source, "python tools/x.py", [("DATA", data), ("N", 3)]))
            target = Path(tmp) / "generated.py"
            target.write_text(text, encoding="utf-8")
            module = load_module(target)
        self.assertEqual((module.DATA, module.N), (data, 3))
        self.assertIn("자동 생성 — 손으로 고치지 말 것", text)
        self.assertIn(_xlsx_common.file_sha256(source) if source.exists() else "", text)


@unittest.skipUnless(CODEBOOK_XLSX and TAXONOMY_XLSX, "원본 xlsx가 없는 환경(저장소 밖 파일)")
class ExtractorReproducibilityTest(unittest.TestCase):
    """원본 xlsx가 있으면: 생성 모듈의 sha256이 원본과 같고, 추출기를 다시 돌린 결과가 커밋된 모듈과 바이트 단위로 같다."""

    def test_codebook_module_is_reproducible(self):
        self.assertEqual(codebook_data.SOURCE_SHA256, _xlsx_common.file_sha256(CODEBOOK_XLSX))
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "codebook_data.py"
            with open(os.devnull, "w") as sink, mock.patch("sys.stdout", sink):
                extract_codebook.main(["--xlsx", CODEBOOK_XLSX, "--out", str(out)])
            self.assertEqual(out.read_bytes(), (RUNNER_DIR / "kyab_runner" / "spec" / "codebook_data.py").read_bytes())

    def test_taxonomy_module_is_reproducible(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "taxonomy_data.py"
            with open(os.devnull, "w") as sink, mock.patch("sys.stdout", sink):
                extract_taxonomy.main(["--xlsx", TAXONOMY_XLSX, "--out", str(out)])
            self.assertEqual(out.read_bytes(), (RUNNER_DIR / "kyab_runner" / "spec" / "taxonomy_data.py").read_bytes())


class ExtractorExitCodeTest(unittest.TestCase):
    """R12: 추출기의 입력·형식 오류(원본 없음·머리글 없음)는 stderr 한 줄 + 종료 2(러너 계약), 선언 불일치만 추출기 고유의 1."""

    def test_missing_xlsx_exits_2(self):
        with tempfile.TemporaryDirectory() as tmp:
            missing = str(Path(tmp) / "none.xlsx")
            for call in (lambda: _xlsx_common.find_xlsx(tmp, "zzz"), lambda: _xlsx_common.resolve_xlsx(missing, "zzz"),
                         lambda: extract_codebook.main(["--xlsx", missing]), lambda: extract_taxonomy.main(["--xlsx", missing])):
                err = io.StringIO()
                with mock.patch("sys.stderr", err), self.assertRaises(SystemExit) as caught:
                    call()
                self.assertEqual(caught.exception.code, _xlsx_common.EXIT_INVALID)
                self.assertIn("xlsx", err.getvalue())

    @staticmethod
    def _workbook(path, declared_by_sheet, field_count):
        """CSV_SHEETS 7장짜리 작은 코드북 xlsx: 제목·'N개 필드' 부제·머리글 행·필드 행 field_count개."""
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for sheet in extract_codebook.CSV_SHEETS:
            ws = wb.create_sheet(sheet)
            ws.append([f"{sheet} 제목"])
            ws.append([f"({declared_by_sheet.get(sheet, field_count)}개 필드)"])
            ws.append([label for label, _ in extract_codebook.HEADER_LABELS])
            for i in range(field_count):
                ws.append(["필수", "", "", f"field_{i}", "설명", "자유 텍스트", "", "", "", ""])
        wb.save(path)

    def test_header_missing_exits_2_and_declared_mismatch_exits_1(self):
        with tempfile.TemporaryDirectory() as tmp:
            out = Path(tmp) / "out.py"
            ok = Path(tmp) / "codebook_ok.xlsx"
            self._workbook(ok, {}, 2)
            with open(os.devnull, "w") as sink, mock.patch("sys.stdout", sink):
                self.assertEqual(extract_codebook.main(["--xlsx", str(ok), "--out", str(out)]), 0)
            self.assertTrue(out.exists())
            out.unlink()
            # 선언 3개 vs 읽은 2개 → 1, 모듈 없음
            mismatch = Path(tmp) / "codebook_mismatch.xlsx"
            self._workbook(mismatch, {"04_runs": 3}, 2)
            err = io.StringIO()
            with open(os.devnull, "w") as sink, mock.patch("sys.stdout", sink), mock.patch("sys.stderr", err):
                self.assertEqual(extract_codebook.main(["--xlsx", str(mismatch), "--out", str(out)]), _xlsx_common.EXIT_MISMATCH)
            self.assertIn("04_runs", err.getvalue())
            self.assertFalse(out.exists())
            # '필드명' 머리글 없음 → 2
            import openpyxl
            wb = openpyxl.load_workbook(ok)
            wb["02_item_tags"].cell(3, 4).value = "이름"
            wb.save(ok)
            err = io.StringIO()
            with open(os.devnull, "w") as sink, mock.patch("sys.stdout", sink), mock.patch("sys.stderr", err), \
                    self.assertRaises(SystemExit) as caught:
                extract_codebook.main(["--xlsx", str(ok), "--out", str(out)])
            self.assertEqual(caught.exception.code, _xlsx_common.EXIT_INVALID)
            self.assertIn("02_item_tags", err.getvalue())
            self.assertFalse(out.exists())


class ExtractorVerifyTest(unittest.TestCase):
    """extract_taxonomy.verify(codebook-04): 앞 검산이 실패해도 예외 없이 전체 목록을 찍고 실패 수를 돌려준다."""

    def test_verify_survives_missing_major(self):
        rows = [{"level_label": "하위", "code": "A1.01", "parent_code": "A1", "parent_name": "x", "name": "n",
                 "definition": "d", "scope_note": "", "source": "", "legacy_code": "R1.01", "source_cells": {}}]
        with open(os.devnull, "w") as sink, mock.patch("sys.stdout", sink):
            failed = extract_taxonomy.verify(rows, [], [])
        self.assertGreater(failed, 0)
