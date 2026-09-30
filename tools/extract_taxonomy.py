"""01_분류팀_벤치마크비교와_재분류.xlsx -> taxonomy_A1-A10.json + crosswalk_RM_to_A.csv

A1~A10 단일 분류체계(상위 10 · 하위 39)를 분류팀 xlsx에서 기계적으로 옮긴다.
extract_codebook.py와 같은 원칙을 따른다.

  * xlsx가 유일한 원천이다. 코드·이름·정의를 이 파일에 손으로 적지 않는다.
  * 셀 원문은 변형하지 않는다(앞뒤 공백만 제거). 각 값에는 출처 셀 주소를 붙인다.
  * 열은 위치가 아니라 머리글 이름으로 찾는다(열이 옮겨져도 깨지지 않게).
  * 같은 내용을 담은 시트끼리 교차 검산하고, 하나라도 어긋나면 파일을 쓰지 않고
    종료 코드 1로 끝낸다.

원천 시트
  코드북49        주 원천. 층위·코드·명칭·정의·포함/경계·출처·이전 코드 (49행)
  분류표_A1-A10   검산용. 대분류별 하위 수·소분류 목록·상위 정의·이전 코드 + 안내 문구
  하위39          검산용. 소분류 39개 목록

실행:  runner/.venv/bin/python runner/tools/extract_taxonomy.py
"""
import csv
import glob
import json
import os
import re
import sys
import unicodedata

import openpyxl

# ── 경로 ────────────────────────────────────────────────────────────────
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
SRC_DIR = os.path.join(ROOT, "project proposal")
SCHEMA_DIR = os.path.join(ROOT, "runner", "schema")
OUT_TAXONOMY = os.path.join(SCHEMA_DIR, "taxonomy_A1-A10.json")
OUT_CROSSWALK = os.path.join(SCHEMA_DIR, "crosswalk_RM_to_A.csv")

# ── 원천 파일·시트 이름 ─────────────────────────────────────────────────
SRC_NAME_KEY = "벤치마크비교와_재분류"   # 파일명에 들어 있는 고정 문자열
SHEET_CODEBOOK = "코드북49"
SHEET_TABLE = "분류표_A1-A10"
SHEET_SUB = "하위39"

# ── 코드 형식 (분류표 A2: "대분류는 A1–A10, 소분류는 A1.01 형식") ────────
RE_MAJOR = r"^A(?:[1-9]|10)$"
RE_SUB = r"^A(?:[1-9]|10)\.[0-9]{2}$"
# 이전 코드 형식: R1~R5 / M01~M05 와 그 하위(.NN)
RE_LEGACY = r"^(R[1-5]|M0[1-5])(\.[0-9]{2})?$"

# ── 기대 개수 (지시문·분류표 기준. 어긋나면 실패 처리) ──────────────────
EXPECT_MAJOR, EXPECT_SUB = 10, 39

CROSSWALK_COLUMNS = ["legacy_code", "legacy_scheme", "new_code", "level",
                     "relation", "new_name", "source_cell"]


# ── 공통 도우미 ─────────────────────────────────────────────────────────
def nfc(text):
    """macOS에서 온 파일명은 자모가 분리(NFD)돼 있어 비교 전에 NFC로 맞춘다."""
    return unicodedata.normalize("NFC", text)


def cell_text(cell):
    """셀 값을 문자열로. 빈 셀은 빈 문자열, 앞뒤 공백만 제거하고 본문은 그대로."""
    return "" if cell.value is None else str(cell.value).strip()


def find_xlsx():
    for path in glob.glob(os.path.join(SRC_DIR, "*.xlsx")):
        if SRC_NAME_KEY in nfc(os.path.basename(path)):
            return path
    sys.exit(f"'{SRC_NAME_KEY}' xlsx를 {SRC_DIR}에서 찾지 못했습니다")


def header_map(ws, header_row, required):
    """머리글 행에서 {머리글: 열 번호}를 만든다. 필요한 머리글이 없으면 중단."""
    found = {cell_text(c): c.column for c in ws[header_row] if cell_text(c)}
    missing = [h for h in required if h not in found]
    if missing:
        sys.exit(f"[{ws.title}] {header_row}행에 머리글 {missing}이 없습니다: {list(found)}")
    return found


def split_code_name(label):
    """'A1 그루밍·성적 착취' -> ('A1', '그루밍·성적 착취')"""
    code, _, name = label.partition(" ")
    return code, name.strip()


class Checker:
    """검산 결과를 모아 출력한다. 하나라도 실패하면 산출물을 쓰지 않는다."""

    def __init__(self):
        self.failed = 0

    def check(self, label, ok, detail=""):
        self.failed += 0 if ok else 1
        print(f"  [{'OK' if ok else 'FAIL'}] {label}" + (f" — {detail}" if detail else ""))


# ── 1. 주 원천: 코드북49 ────────────────────────────────────────────────
def read_codebook49(ws):
    """코드북49의 49행을 시트 순서 그대로 읽는다.

    반환: 행 dict 목록. 상위 행은 parent_code가 빈 문자열이다.
    """
    cols = header_map(ws, 1, ["층위", "코드", "대분류 코드와 이름", "명칭",
                              "정의", "포함·경계", "출처", "이전 코드"])
    rows = []
    for r in range(2, ws.max_row + 1):
        get = lambda h: ws.cell(r, cols[h])          # noqa: E731 (행 고정 접근자)
        code = cell_text(get("코드"))
        if not code:
            continue
        parent_code, parent_name = split_code_name(cell_text(get("대분류 코드와 이름")))
        rows.append({
            "level_label": cell_text(get("층위")),   # 원문 값: '상위' | '하위'
            "code": code,
            "parent_code": parent_code,
            "parent_name": parent_name,
            "name": cell_text(get("명칭")),
            "definition": cell_text(get("정의")),
            "scope_note": cell_text(get("포함·경계")),
            "source": cell_text(get("출처")),
            "legacy_code": cell_text(get("이전 코드")),
            "source_cells": {
                "code": f"{ws.title}!{get('코드').coordinate}",
                "legacy_code": f"{ws.title}!{get('이전 코드').coordinate}",
            },
        })
    return rows


# ── 2. 검산 원천: 분류표_A1-A10 ─────────────────────────────────────────
def read_overview(ws):
    """분류표 시트에서 대분류 요약 행과 안내 문구를 읽는다.

    반환: (대분류 요약 목록, 안내 문구 목록[{cell, text}])
    """
    # 머리글 행은 'A열 = 코드'인 행이다(위쪽 안내 문구 길이가 바뀌어도 찾도록).
    header_row = next((r for r in range(1, ws.max_row + 1)
                       if cell_text(ws.cell(r, 1)) == "코드"), None)
    if header_row is None:
        sys.exit(f"[{ws.title}] '코드' 머리글 행을 찾지 못했습니다")
    cols = header_map(ws, header_row, ["코드", "대분류 코드와 이름", "하위 수",
                                       "소분류 코드와 이름", "상위 정의", "이전 코드"])
    majors, notes = [], []
    for r in range(1, ws.max_row + 1):
        first = cell_text(ws.cell(r, 1))
        if not first or r == header_row:
            continue
        if re.match(RE_MAJOR, first):
            sub_labels = cell_text(ws.cell(r, cols["소분류 코드와 이름"])).split("\n")
            majors.append({
                "code": first,
                "name": split_code_name(cell_text(ws.cell(r, cols["대분류 코드와 이름"])))[1],
                "sub_count": int(cell_text(ws.cell(r, cols["하위 수"]))),
                "subs": [split_code_name(s.strip()) for s in sub_labels if s.strip()],
                "definition": cell_text(ws.cell(r, cols["상위 정의"])),
                "legacy_code": cell_text(ws.cell(r, cols["이전 코드"])),
            })
        else:
            # 제목·안내·출처 문구. 해석하지 않고 원문 그대로 보존한다.
            notes.append({"cell": f"{ws.title}!A{r}", "text": first})
    return majors, notes


def read_sub39(ws):
    """하위39 시트의 (소분류 코드, 이름, 대분류 코드) 목록."""
    cols = header_map(ws, 1, ["소분류 코드와 이름", "대분류 코드와 이름"])
    out = []
    for r in range(2, ws.max_row + 1):
        label = cell_text(ws.cell(r, cols["소분류 코드와 이름"]))
        if label:
            code, name = split_code_name(label)
            parent = split_code_name(cell_text(ws.cell(r, cols["대분류 코드와 이름"])))[0]
            out.append((code, name, parent))
    return out


# ── 3. 검산 ─────────────────────────────────────────────────────────────
def verify(rows, overview, sub39):
    """코드북49 내부 정합성 + 분류표·하위39와의 일치를 확인한다."""
    ck = Checker()
    majors = [r for r in rows if r["level_label"] == "상위"]
    subs = [r for r in rows if r["level_label"] == "하위"]
    major_codes = [m["code"] for m in majors]

    print("코드북49 내부 검산")
    ck.check("층위 값은 '상위'·'하위'뿐", len(majors) + len(subs) == len(rows),
             f"전체 {len(rows)}행")
    ck.check(f"상위 {EXPECT_MAJOR}개", len(majors) == EXPECT_MAJOR, f"실제 {len(majors)}")
    ck.check(f"하위 {EXPECT_SUB}개", len(subs) == EXPECT_SUB, f"실제 {len(subs)}")
    ck.check("코드 중복 없음", len({r["code"] for r in rows}) == len(rows))
    ck.check("상위 코드 형식 A1~A10", all(re.match(RE_MAJOR, c) for c in major_codes))
    ck.check("하위 코드 형식 A#.##", all(re.match(RE_SUB, s["code"]) for s in subs))
    ck.check("하위의 대분류가 상위 목록에 존재", all(s["parent_code"] in major_codes for s in subs))
    ck.check("하위 코드 접두부 = 대분류 코드",
             all(s["code"].split(".")[0] == s["parent_code"] for s in subs))
    ck.check("하위 행의 대분류 이름 = 상위 행 명칭",
             all(s["parent_name"] == next(m["name"] for m in majors if m["code"] == s["parent_code"])
                 for s in subs))
    ck.check("명칭·정의 빈칸 없음", all(r["name"] and r["definition"] for r in rows))

    print("이전 코드(코드북49 '이전 코드' 열) 검산")
    legacy = [r["legacy_code"] for r in rows]
    ck.check("이전 코드 빈칸 없음·형식 R#/M0#(.##)", all(re.match(RE_LEGACY, c) for c in legacy))
    ck.check("이전 코드 중복 없음(1:1)", len(set(legacy)) == len(legacy))
    ck.check("층위 보존(상위↔상위, 하위↔하위)",
             all(("." in r["legacy_code"]) == (r["level_label"] == "하위") for r in rows))
    legacy_of = {r["code"]: r["legacy_code"] for r in rows}
    ck.check("하위의 이전 코드 부모 = 대분류의 이전 코드, 하위 번호 보존",
             all(s["legacy_code"] == f"{legacy_of[s['parent_code']]}.{s['code'].split('.')[1]}"
                 for s in subs))

    print("분류표_A1-A10 대조")
    ov = {m["code"]: m for m in overview}
    ck.check("대분류 코드 집합·순서 일치", [m["code"] for m in overview] == major_codes)
    ck.check("대분류 명칭 일치", all(ov[m["code"]]["name"] == m["name"] for m in majors))
    ck.check("상위 정의 일치", all(ov[m["code"]]["definition"] == m["definition"] for m in majors))
    ck.check("대분류 이전 코드 일치", all(ov[m["code"]]["legacy_code"] == m["legacy_code"] for m in majors))
    subs_by_parent = {c: [(s["code"], s["name"]) for s in subs if s["parent_code"] == c]
                      for c in major_codes}
    ck.check("하위 수(C열) = 실제 하위 개수",
             all(ov[c]["sub_count"] == len(subs_by_parent[c]) for c in major_codes),
             ", ".join(f"{c}:{len(subs_by_parent[c])}" for c in major_codes))
    ck.check("소분류 목록(D열) 코드·이름·순서 일치",
             all(ov[c]["subs"] == subs_by_parent[c] for c in major_codes))

    print("하위39 대조")
    ck.check("소분류 코드·이름·대분류·순서 일치",
             sub39 == [(s["code"], s["name"], s["parent_code"]) for s in subs],
             f"하위39 {len(sub39)}행")
    return ck.failed


# ── 4. 산출물 구성 ──────────────────────────────────────────────────────
def legacy_scheme(code):
    """이전 코드의 계열. R = 기존 위험군, M = 기존 경계 검토 영역."""
    return code[0]


def build_taxonomy(src_path, rows, notes):
    """검증용 JSON. 대분류 아래에 소분류를 시트 순서대로 넣는다."""
    def entry(r):
        return {"code": r["code"], "name": r["name"], "definition": r["definition"],
                "scope_note": r["scope_note"], "source": r["source"],
                "legacy_code": r["legacy_code"], "source_cells": r["source_cells"]}

    majors = []
    for r in rows:
        if r["level_label"] == "상위":
            majors.append({**entry(r), "sort_order": len(majors) + 1, "subs": []})
        else:
            parent = next(m for m in majors if m["code"] == r["parent_code"])
            parent["subs"].append({**entry(r), "parent_code": r["parent_code"],
                                   "sort_order": len(parent["subs"]) + 1})
    return {
        "source_file": nfc(os.path.basename(src_path)),
        "source_sheets": [SHEET_CODEBOOK, SHEET_TABLE, SHEET_SUB],
        # 분류체계 판본은 xlsx에 적혀 있지 않다. 확정되면 값만 채운다(임의 기입 금지).
        "taxonomy_version": None,
        "code_format": {"major_regex": RE_MAJOR, "sub_regex": RE_SUB},
        "counts": {"major": len(majors), "sub": sum(len(m["subs"]) for m in majors)},
        # 분류표 시트의 제목·안내·출처 문구 원문. 확정 범위 해석의 근거로 남긴다.
        "source_notes": notes,
        "majors": majors,
    }


def build_crosswalk(rows):
    """이전 코드 -> 새 코드 대응 행.

    relation 값은 xlsx에서 기계적으로 정해진다.
      code_rename            R 계열. 코드 표기만 바뀐다.
      code_rename_promoted   M 계열. 코드 표기가 바뀌고, '경계 검토 영역'에서
                             주 분류에 쓰는 대분류로 지위가 바뀐다.
    """
    out = []
    for r in rows:
        scheme = legacy_scheme(r["legacy_code"])
        out.append({
            "legacy_code": r["legacy_code"],
            "legacy_scheme": scheme,
            "new_code": r["code"],
            "level": "major" if r["level_label"] == "상위" else "sub",
            "relation": "code_rename" if scheme == "R" else "code_rename_promoted",
            "new_name": r["name"],
            "source_cell": r["source_cells"]["legacy_code"],
        })
    return out


def main():
    src = find_xlsx()
    # data_only=True: 수식이 있어도 저장된 값을 읽는다. 원본은 읽기만 한다.
    wb = openpyxl.load_workbook(src, data_only=True)
    rows = read_codebook49(wb[SHEET_CODEBOOK])
    overview, notes = read_overview(wb[SHEET_TABLE])
    sub39 = read_sub39(wb[SHEET_SUB])

    failed = verify(rows, overview, sub39)
    if failed:
        sys.exit(f"검산 실패 {failed}건 — 산출물을 쓰지 않았습니다")

    os.makedirs(SCHEMA_DIR, exist_ok=True)
    taxonomy = build_taxonomy(src, rows, notes)
    with open(OUT_TAXONOMY, "w", encoding="utf-8") as f:
        json.dump(taxonomy, f, ensure_ascii=False, indent=2)
    crosswalk = build_crosswalk(rows)
    with open(OUT_CROSSWALK, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=CROSSWALK_COLUMNS)
        writer.writeheader()
        writer.writerows(crosswalk)

    by_relation = {}
    for row in crosswalk:
        by_relation[row["relation"]] = by_relation.get(row["relation"], 0) + 1
    print(f"wrote {OUT_TAXONOMY}  (상위 {taxonomy['counts']['major']} · 하위 {taxonomy['counts']['sub']})")
    print(f"wrote {OUT_CROSSWALK}  ({len(crosswalk)}행: {by_relation})")


if __name__ == "__main__":
    # openpyxl이 '데이터 유효성 확장 미지원' 경고를 내지만 값 읽기에는 영향이 없다.
    import warnings
    warnings.filterwarnings("ignore", message="Data Validation extension")
    main()
