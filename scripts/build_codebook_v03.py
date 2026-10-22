"""Codebook v0.3 (draft) as implemented by the runner: docs/codebook_v0.2.xlsx + the changes in
CODEBOOK_V03_CHANGES.md. Same sheets, columns and styling as v0.2; every changed cell is filled yellow, new
rows are marked in column J, and 08 gets one change-history row per change.

    .venv/bin/python scripts/build_codebook_v03.py [--out reports/codebook_v0.3_draft_2026-10-22.xlsx]

tests/test_codebook_v03.py checks that the field lists of the generated file equal kyab.schemas.
"""
from __future__ import annotations

import argparse
import copy
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from kyab.schemas import PERSONA_PROMPTS, PERSONAS, TABLES  # noqa: E402

SRC = ROOT / "docs" / "codebook_v0.2.xlsx"
YELLOW = PatternFill("solid", fgColor="FFF2CC")
TODAY = date(2026, 10, 22)
TAG = "v0.3(안)"
COLS = "ABCDEFGHIJ"  # 생성 단계·필수성, AI 전달, 입력 주체, 필드명, 필드 설명, 값·형식, 필요한 이유, 키·연결, 운영 메모, 근거·합의

# (sheet, field) → {column letter: new text}
EDITS: dict[tuple[str, str], dict[str, str]] = {
    ("01_items", "protocol_id"): {
        "F": "통제어휘: ST1-1.0.0, MT3-1.0.0, MT7-1.0.0, MT10-1.0.0 (구 ST1-v1 등은 과거 기록 판독만)",
        "J": "09/30 회신 확정"},
    ("01_items", "original_text"): {
        "A": "조건부", "F": "원문 잠금본. source_benchmark=NEW(신규 작성)가 아니면 필수", "J": "09/30 회신 확정"},
    ("01_items", "target_age_group"): {
        "I": "페르소나 축으로 실행하면 응답의 연령대는 04 persona_id의 페르소나(09 시트)가 기준. 이 열은 문항 설계값"},
    ("01_items", "user_gender"): {
        "I": "페르소나 축으로 실행하면 응답의 성별은 04 persona_id의 페르소나(09 시트)가 기준. 이 열은 문항 설계값"},
    ("02_item_tags", "taxonomy_version"): {
        "F": "MAJOR.MINOR.PATCH. MAJOR 0 = 이전 R/M 체계(예: 0.7.0), MAJOR 1 이상 = A1~A10 체계(현재 1.0.0)",
        "J": "09/30 회신(R/M → A1~A10 통합) · A1~A10 확정(10월 회의)"},
    ("02_item_tags", "primary_risk"): {
        "F": "taxonomy_version MAJOR ≥ 1: A1~A10 중 1개 / MAJOR 0: R1~R5 중 1개. "
             "risk_review_status=mapped일 때만 필수, not_applicable(대조 문항)이면 공란",
        "J": "09/30 회신 확정"},
    ("02_item_tags", "sub_risk_codes"): {
        "F": "taxonomy_version MAJOR ≥ 1: 소분류 코드(A1.01 꼴, 39개)의 JSON 배열. 각 코드는 primary_risk의 하위여야 함",
        "J": "09/30 회신 확정 · 개수 제한(1개)은 미정"},
    ("02_item_tags", "secondary_risks"): {
        "F": "primary_risk와 같은 체계의 대분류 코드 JSON 배열 (MAJOR ≥ 1: A1~A10, 0: R1~R5)",
        "J": "09/30 회신 확정"},
    ("02_item_tags", "risk_review_status"): {
        "F": "통제어휘: unreviewed, mapped, hold, scope_review, not_applicable (not_applicable = 대조 문항)",
        "J": "09/30 회신 확정"},
    ("02_item_tags", "m_review_codes"): {
        "I": "A1~A10 체계(taxonomy_version MAJOR ≥ 1)에서는 M01~M05가 A6~A10에 통합되어 신규 입력 불필요. 이전 기록 보존용"},
    ("02_item_tags", "intent_tag"): {"A": "선택", "J": "09/30 회신 확정"},
    ("04_runs", "rollout_no"): {
        "F": "정수 1 이상. 과업 보고는 1~3(3회), 추가 반복은 4부터 이어 붙임(공개 모델 추가 반복, 논문용 100회 이상)",
        "J": "10월 회의 결정"},
    ("04_runs", "protocol_id"): {
        "F": "통제어휘: ST1-1.0.0, MT3-1.0.0, MT7-1.0.0, MT10-1.0.0 (구 -v1은 과거 기록 판독만)",
        "J": "09/30 회신 확정"},
    ("04_runs", "temperature"): {
        "F": "숫자 [0, 2]. 실행에 쓴 값을 그대로 기록. 본평가 1.0 (반복 실행 · GPT-5 mini는 1만 허용)",
        "I": "0.0이면 반복해도 같은 응답이 나와 반복의 의미가 없음",
        "J": "10월 회의 결정(반복 3회) · 값 1.0은 연구실 A 제안"},
    ("04_runs", "max_output_tokens"): {
        "F": "양의 정수 8192 (v0.2 시절 기록은 1024)",
        "I": "3턴 대화는 서버 최대 문맥 길이 32,768 필요(앞 턴 응답 누적)",
        "J": "10/05 회신 ① 출력 한도 상향 · 값 8,192는 연구실 A 결정"},
}

# (sheet, insert after field) → new row values A..J
NEW_ROWS: dict[tuple[str, str], list[str]] = {
    ("04_runs", "rollout_no"): [
        "실행 자동기록 필수", "×", "시스템 생성", "persona_id",
        "이 실행에 적용한 페르소나 ID입니다. 페르소나 축 없이 실행하면 none.",
        "정규식 ^(none|PERSONA-[A-Z0-9][A-Z0-9_-]*)$ 예: PERSONA-MID-F",
        "같은 문항을 연령대 3 × 성별 2 페르소나로 실행한 결과를 구분·집계한다.",
        "FK → personas.persona_id (none 제외). 실행 단위 키: 모델 × 데이터셋 × 문항 × protocol × persona × rollout",
        "모델에는 전달하지 않음(페르소나별로 다시 쓴 사용자 발화만 전달)",
        f"{TAG} 신규 · 10월 회의 결정(페르소나 6벌)"],
    ("02_item_tags", "secondary_risks"): [
        "조건부", "×", "사람 입력·검토", "control_target_risk",
        "대조 문항(case_type=safe_control)이 어느 위험군의 대조인지 나타냅니다.",
        "taxonomy_version MAJOR ≥ 1: A1~A10 중 1개 / MAJOR 0: R1~R5. 대조 문항이 아니면 공란",
        "위험군별 과잉거절률(ORR)을 산출한다.",
        "→ taxonomy 기준표 대분류",
        "열 이름·위치는 잠정. 실행 러너는 이 열이 없는 파일도 읽음",
        f"{TAG} 신규 · 10/05 회신 ② 확정(열 추가) · 이름·규칙 잠정"],
}

HISTORY = [
    ("01·04 protocol_id", "ST1-v1 등", "ST1-1.0.0, MT3-1.0.0, MT7-1.0.0, MT10-1.0.0", "판본 표기 SemVer 통일", "09/30 회신"),
    ("02 intent_tag · 01 original_text", "필수", "intent_tag 선택 / original_text 조건부(NEW 외 필수)", "입력 부담·신규 문항 처리", "09/30 회신"),
    ("02 primary_risk·secondary_risks·sub_risk_codes·taxonomy_version",
     "R1~R5 (+M01~M05)", "A1~A10·소분류 39 (taxonomy_version 1.0.0), MAJOR 0은 이전 R/M", "위험 분류 통합", "09/30 회신 · 10월 회의"),
    ("02 risk_review_status", "4개 값", "not_applicable 추가, primary_risk는 mapped일 때만 필수", "대조 문항 처리", "09/30 회신"),
    ("02 control_target_risk (신규)", "없음", "대조 문항의 대상 위험군 1개", "위험군별 과잉거절률", "10/05 회신 ② (이름 잠정)"),
    ("04 max_output_tokens", "1024 고정", "8192", "긴 응답이 잘려 상담 안내 누락으로 판정되는 것 방지", "10/05 회신 ①"),
    ("06 critical_failure_code", "미발생 표기 미정", "NONE", "빈칸(미채점)과 구분", "10/05 회신 ③"),
    ("06·07 inconclusive", "허용 여부 미정", "허용. 실패율에서 빼면 보류 건수·비율·분모를 함께 제시", "보류가 많은 모델이 좋아 보이는 것 방지",
     "10/05 회신 ④"),
    ("04 persona_id (신규)", "없음", "페르소나 ID (연령대 3 × 성별 2 = 6벌), 없으면 none", "같은 문항의 페르소나별 실행 구분", "10월 회의"),
    ("04 rollout_no", "1, 2, 3", "1 이상 정수 (보고 3회, 추가 반복은 4부터)", "공개 모델 추가 반복·논문용 반복", "10월 회의"),
    ("04 temperature", "0.0 고정", "기록되는 설정값 [0, 2], 본평가 1.0", "반복 3회 결정(0.0이면 반복 무의미), GPT-5 mini는 1만 허용",
     "10월 회의 · 값은 연구실 A 제안"),
    ("09 personas · persona_prompts (확장)", "없음", "페르소나 기준표와 페르소나별 재작성 발화 CSV", "페르소나 축 실행 입력", "10월 회의"),
]


def find_row(ws, field: str) -> int:
    for r in range(4, ws.max_row + 1):
        if ws.cell(r, 4).value == field:
            return r
    raise KeyError(f"{ws.title}: {field}")


def copy_row_style(ws, src: int, dst: int, ncol: int = 10) -> None:
    for c in range(1, ncol + 1):
        s, d = ws.cell(src, c), ws.cell(dst, c)
        d._style = copy.copy(s._style)
    ws.row_dimensions[dst].height = ws.row_dimensions[src].height


def insert_rows_keep_merges(ws, idx: int, amount: int) -> None:
    """openpyxl's insert_rows does not move merged ranges; shift the ones at or below idx by hand."""
    below = [m for m in ws.merged_cells.ranges if m.min_row >= idx]
    for m in below:
        ws.unmerge_cells(str(m))
    ws.insert_rows(idx, amount)
    for m in below:
        ws.merge_cells(start_row=m.min_row + amount, start_column=m.min_col,
                       end_row=m.max_row + amount, end_column=m.max_col)


def field_count(ws) -> int:
    return sum(1 for r in range(4, ws.max_row + 1) if ws.cell(r, 4).value)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=str(ROOT / "reports" / "codebook_v0.3_draft_2026-10-22.xlsx"))
    args = ap.parse_args(argv)
    wb = openpyxl.load_workbook(SRC)

    for (sheet, field), cols in EDITS.items():
        ws = wb[sheet]
        r = find_row(ws, field)
        for col, text in cols.items():
            cell = ws[f"{col}{r}"]
            cell.value = text
            cell.fill = YELLOW
    for (sheet, after), values in NEW_ROWS.items():
        ws = wb[sheet]
        r = find_row(ws, after) + 1
        insert_rows_keep_merges(ws, r, 1)
        copy_row_style(ws, r - 1, r)
        for col, v in zip(COLS, values):
            ws[f"{col}{r}"].value = v
            ws[f"{col}{r}"].fill = YELLOW
    for name, t in TABLES.items():
        ws = wb[t.sheet]
        sub = ws["A2"].value or ""
        n = field_count(ws)
        rest = sub.split("·", 2)[2] if sub.count("·") >= 2 else sub
        ws["A2"].value = f"{TAG} 기준 · {n}개 필드 ·{rest}"

    ov = wb["00_관계·운영원칙"]
    ov["A2"].value = f"청소년 AI 위험 벤치마크 · 7개 CSV 코드북 ({TAG}, 실행 러너 runner-v0.2.1 구현 기준, {TODAY.isoformat()})"
    for row in ov.iter_rows(min_row=3, max_row=3):
        for c in row:
            if c.value == "구조 협의용 초안":
                c.value = "v0.3(안) · 노란 칸 = v0.2 대비 변경 · 승인 전"
                c.fill = YELLOW

    # 08: change history rows after the existing ones, open-issue statuses updated
    h = wb["08_변경이력·미확정"]
    header = next(r for r in range(1, h.max_row + 1) if h.cell(r, 1).value == "변경일")
    last = header
    while h.cell(last + 1, 2).value:
        last += 1
    insert_rows_keep_merges(h, last + 1, len(HISTORY))
    for i, (target, before, after, why, basis) in enumerate(HISTORY, 1):
        r = last + i
        copy_row_style(h, last, r, 8)
        for c, v in enumerate([TODAY, target, before, after, why, f"코드북 {TAG}", f"{basis} · v0.3 반영 요청",
                               "실행 러너 runner-v0.2.1에 반영"], 1):
            h.cell(r, c).value = v
            h.cell(r, c).fill = YELLOW
        h.cell(r, 1).number_format = h.cell(last, 1).number_format
    status = {"R/M 위험군·태그 허용값": "A1~A10·소분류 39로 통합 확정 (taxonomy_version 1.0.0). L1 상위층은 전문가 미팅 후",
              "중대 실패 코드 미발생 기록": "확정: NONE (10/05 회신 ③)",
              "최종 판정값 허용 범위": "확정: inconclusive 허용, 실패율에서 빼면 보류 건수·비율·분모 함께 제시 (10/05 회신 ④)"}
    for r in range(1, h.max_row + 1):
        k = h.cell(r, 1).value
        if k in status:
            h.cell(r, 6).value = status[k]
            h.cell(r, 6).fill = YELLOW

    # 09: extension tables for the persona axis (same column layout as the table sheets)
    src = wb["03_prompts"]
    ps = wb.copy_worksheet(src)
    ps.title = "09_personas(확장)"
    for r in range(4, ps.max_row + 1):
        for c in range(1, 11):
            ps.cell(r, c).value = None
    ps["A1"].value = "09 페르소나 확장 CSV 코드북 (personas.csv · persona_prompts.csv)"
    ps["A2"].value = (f"{TAG} 신규 · 7개 CSV 밖의 확장 입력 · 같은 문항을 연령대 3 × 성별 2 = 6개 페르소나로 실행. "
                      "모델에는 persona_prompts.message_text만 전달")
    rows = [
        ["personas.csv", None, None, None, None, None, None, None, None, None],
        ["문항·태그 입력 필수", "×", "페르소나 담당", "persona_id", "페르소나 고유 ID", "정규식 ^PERSONA-[A-Z0-9][A-Z0-9_-]*$ 예: PERSONA-MID-F",
         "실행(04 persona_id)과 발화를 연결", "PK", None, "10월 회의 결정"],
        ["문항·태그 입력 필수", "×", "페르소나 담당", "persona_version", "페르소나 정의의 판본", "SemVer", "정의가 바뀌면 이전 결과와 구분",
         None, None, "연구실 A 제안"],
        ["문항·태그 입력 필수", "×", "페르소나 담당", "age_band", "연령대", "통제어휘: early_10_12, mid_13_15, late_16_18",
         "연령대별 집계", None, "6벌 = 연령대 3 × 성별 2. 조합이 빠지면 실행 중단", "10월 회의 결정"],
        ["문항·태그 입력 필수", "×", "페르소나 담당", "gender", "성별", "통제어휘: male, female", "성별 집계", None, None, "10월 회의 결정"],
        ["조건부", "×", "페르소나 담당", "persona_label", "사람이 읽는 이름", "자유 텍스트 예: 중학생(13~15세) 여학생", "보고서 표기", None,
         "모델에 전달하지 않음", "연구실 A 제안"],
        ["조건부", "×", "페르소나 담당", "style_note", "발화 스타일 설명", "자유 텍스트", "재작성 기준 기록", None, "모델에 전달하지 않음",
         "연구실 A 제안"],
        ["문항·태그 입력 필수", "×", "페르소나 담당", "persona_status", "자료 상태", "통제어휘: placeholder, provided",
         "자리표시 실행과 실제 페르소나 실행을 구분", None, "placeholder는 개발용", "연구실 A 제안"],
        [None] * 10,
        ["persona_prompts.csv", None, None, None, None, None, None, None, None, None],
        ["문항·태그 입력 필수", "×", "페르소나 담당", "persona_id", "페르소나 ID", "personas.persona_id 값", "발화의 페르소나",
         "복합 PK: persona_id + turn_id · FK → personas", None, "10월 회의 결정"],
        ["문항·태그 입력 필수", "×", "페르소나 담당", "turn_id", "원본 사용자 발화 ID", "03 prompts의 turn_id", "어느 발화를 다시 썼는지",
         "복합 PK 구성 · FK → prompts", "실행 대상인 모든 (문항, 페르소나, turn) 조합이 있어야 함. 하나라도 빠지면 실행 중단", "10월 회의 결정"],
        ["문항·태그 입력 필수", "○", "페르소나 담당", "message_text", "그 페르소나의 말투로 다시 쓴 사용자 발화",
         "한국어 UTF-8(NFC) 텍스트", "같은 문항을 페르소나별 발화로 실행", None, "모델에 전달되는 유일한 페르소나 정보", "10월 회의 결정"],
    ]
    for i, vals in enumerate(rows, 4):
        if i > src.max_row:
            copy_row_style(ps, 4, i)
        for c, v in enumerate(vals, 1):
            ps.cell(i, c).value = v
            ps.cell(i, c).fill = YELLOW if v else PatternFill()

    out = Path(args.out)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
