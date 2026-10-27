"""Readable results workbook for the research team (requested: full dialogue + conditions and
rationale + rubric/CRRI scores + final verdict, hard cases separately).

Called from package.py; reads the per-dialogue bundles already written under report/bundles.
"""
import json

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

HEAD = Font(bold=True)
FILL = PatternFill('solid', fgColor='E3EFEB')
VERDICT_FILL = {'fail': PatternFill('solid', fgColor='F6D5D1'), 'inconclusive': PatternFill('solid', fgColor='F7EBC8'),
                'pass': PatternFill('solid', fgColor='DCEFE0')}
WRAP = Alignment(wrap_text=True, vertical='top')
MAX_CELL = 32000  # Excel cell limit is 32,767 characters


def _cell(v):
    if isinstance(v, (dict, list)):
        v = json.dumps(v, ensure_ascii=False)
    if isinstance(v, str) and len(v) > MAX_CELL:
        v = v[:MAX_CELL] + ' …(잘림: 원문은 bundles JSON)'
    return v


def _sheet(wb, title, header, rows, widths, verdict_col=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill, c.alignment = HEAD, FILL, WRAP
    for r in rows:
        ws.append([_cell(v) for v in r])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
        if verdict_col is not None:
            f = VERDICT_FILL.get(row[verdict_col].value)
            if f:
                row[verdict_col].fill = f
    for i, w in enumerate(widths):
        ws.column_dimensions[ws.cell(1, i + 1).column_letter].width = w
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    return ws


def _ev(evidence):
    return '\n'.join(f"[{e['message_id']}] {e['quote']}" for e in evidence)


def build(bundles, cases, pooled, path, meta):
    wb = Workbook()
    ws = wb.active
    ws.title = '00_안내'
    for line in [f"실행: {meta['run_tag']}  ·  조건표: {meta['card_file']}  ·  smoke(점검용): {meta['smoke']}",
                 '시트: 01_결과요약(대화×채점 모델 1행) / 02_대화원문 / 03_턴별점수 / 04_E·F·CFC / 05_적용조건·근거 / 06_집계 / 07_판단어려운사례',
                 '판정 색: 빨강 fail · 노랑 inconclusive · 초록 pass. 원평정 원본은 judgments/, 계산 결과는 results/에 있음.',
                 '채점 모델별로 결과를 따로 본다(섞어서 집계하지 않음). "자기 채점"은 평가 대상과 채점 모델이 같은 경우.']:
        ws.append([line])
    ws.column_dimensions['A'].width = 140

    rows = []
    for b in bundles:
        r, d, pl = b['result'], b['dialogue'], b['result']['_pipeline']
        rows.append([pl['judge'], pl['target'], r['run']['item_id'], d['condition']['age_band'], d['condition']['gender'],
                     r['safety_verdict'], r['cfc_status'], r['first_failure_turn'],
                     '; '.join(f"{e['criterion_id']}@T{e['turn']}" for e in r['failure_evidence']),
                     r['MRS_app'], r['MRS6'], r['CRRI']['state'], r['CRRI']['value'], r['CRRI']['max_observed_turn_risk'],
                     r['over_refusal'], r['referral'], len(r['human_review']), 'Y' if pl['self_judged'] else '',
                     r['run']['execution_status'], r['run']['completion'], '\n'.join(pl['judge_notes'] + pl['generation_flags'])])
    _sheet(wb, '01_결과요약', ['채점 모델', '평가 대상', '문항', '연령', '성별', '판정', 'CFC', '첫 실패 턴', '실패 근거', 'MRS_app', 'MRS6',
                            'CRRI 상태', 'CRRI', '최대 턴 위험', '과잉거절', '연결', '인간검토 수', '자기 채점', '실행 상태', '완결', '비고'],
           rows, [11, 11, 20, 7, 5, 11, 9, 8, 22, 8, 7, 18, 7, 8, 8, 8, 8, 7, 8, 9, 50], verdict_col=5)

    rows, seen = [], set()
    for b in bundles:
        d = b['dialogue']
        if d['run_id'] in seen:
            continue
        seen.add(d['run_id'])
        for m in d['messages']:
            rows.append([d['target']['name'], d['item_id'], d['condition']['age_band'], d['condition']['gender'],
                         m['turn'], m['role'], m['integrity'], m['content']])
    _sheet(wb, '02_대화원문', ['평가 대상', '문항', '연령', '성별', '턴', '역할', '무결성', '내용'], rows, [11, 20, 7, 5, 5, 9, 9, 110])

    rows = []
    for b in bundles:
        pl = b['result']['_pipeline']
        for j in b['judgments']:
            for s in j['scores']:
                rows.append([pl['judge'], pl['target'], b['result']['run']['item_id'], pl['condition']['age_band'], pl['condition']['gender'],
                             j['phase'], j['target_turn'], s['criterion_id'], s['state'], s['score'], s['anchor_id'], s['reason'], _ev(s['evidence'])])
    _sheet(wb, '03_턴별점수', ['채점 모델', '평가 대상', '문항', '연령', '성별', '단계', '턴', '기준', '상태', '점수', '앵커', '이유', '근거 인용'],
           rows, [11, 11, 20, 7, 5, 8, 5, 6, 12, 6, 8, 60, 60])

    rows = []
    for b in bundles:
        pl = b['result']['_pipeline']
        texts = {x['id']: x['text'] for p in b['conditions']['turn_plan'] for x in p['essential_actions'] + p['failure_conditions']}
        texts.update({x['id']: x['text'] for x in b['conditions']['dialogue_essential_actions'] + b['conditions']['dialogue_failure_conditions']})
        for j in b['judgments']:
            for kind in ['essential_actions', 'failure_conditions', 'cfc_checks', 'a5_components']:
                for x in j[kind]:
                    rows.append([pl['judge'], pl['target'], b['result']['run']['item_id'], pl['condition']['age_band'], pl['condition']['gender'],
                                 j['phase'], j['target_turn'], kind, x['id'], texts.get(x['id'], ''), x['status'], x['reason'], _ev(x['evidence'])])
            for kind in ['over_refusal', 'referral']:
                x = j[kind]
                rows.append([pl['judge'], pl['target'], b['result']['run']['item_id'], pl['condition']['age_band'], pl['condition']['gender'],
                             j['phase'], j['target_turn'], kind, kind, '', x['status'], x['reason'], _ev(x['evidence'])])
    _sheet(wb, '04_E·F·CFC', ['채점 모델', '평가 대상', '문항', '연령', '성별', '단계', '턴', '종류', 'ID', '조건 문구', '상태', '이유', '근거 인용'],
           rows, [11, 11, 20, 7, 5, 8, 5, 16, 18, 45, 11, 50, 50])

    rows, seen = [], set()
    for b in bundles:
        iid = b['result']['run']['item_id']
        if iid in seen:
            continue
        seen.add(iid)
        why = {t['turn']: t for t in b['rationale']['turns']}
        for p in b['conditions']['turn_plan']:
            w = why[p['turn']]
            rows.append([iid, p['turn'], ', '.join(p['applied_dimensions']), ', '.join(p['required_dimensions']),
                         '\n'.join(f"{x['id']} [{x['dimension']}] {x['text']}" for x in p['essential_actions']),
                         '\n'.join(f"{x['id']} [{x['dimension']}] {x['text']}" for x in p['failure_conditions']),
                         'Y' if p['referral_required'] else '', '\n'.join(p['referral_criteria']),
                         f"[적용] {w['applied']}\n[필수] {w['required']}\n[D4] {w['d4']}\n[연결] {w['referral']}",
                         '\n'.join(f'{k}: {v}' for k, v in w['conditions'].items()), '\n'.join(b['card_flags'])])
    _sheet(wb, '05_적용조건·근거', ['문항', '턴', '적용 차원', '필수 차원', '필수 행동 E', '일반 실패 F', '연결 필수', '연결 기준', '적용·필수 근거', 'E/F 근거', '플래그'],
           rows, [20, 5, 16, 12, 45, 45, 8, 40, 60, 60, 50])

    rows = []
    for key, rates in pooled.items():
        for m, x in rates.items():
            rows.append([key, m, x['numerator'], x['denominator'], x['value']])
    _sheet(wb, '06_집계', ['채점 모델 | 평가 대상', '지표', '분자', '분모', '값'], rows, [30, 18, 8, 8, 10])

    rows = [[c['judge'], c['target'], c['item_id'], c['condition']['age_band'], c['condition']['gender'], c['verdict'], '\n'.join(c['reasons'])] for c in cases]
    _sheet(wb, '07_판단어려운사례', ['채점 모델', '평가 대상', '문항', '연령', '성별', '판정', '사유'], rows, [14, 11, 20, 7, 5, 22, 90])
    wb.save(path)
