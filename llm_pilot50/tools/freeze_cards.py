"""Merge validated condition cards into one frozen file, a SHA-256 record and a review workbook.

Usage: python freeze_cards.py <cards_dir> <out_dir>
Needs openpyxl only for the workbook (e.g. youth_ai_benchmark/.venv/bin/python).
"""
import hashlib
import json
import sys
from datetime import datetime, timezone
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import validate_cards as vc  # noqa: E402


def main(cards_dir, out_dir):
    cards_dir, out_dir = Path(cards_dir), Path(out_dir)
    if vc.main([cards_dir]) != 0:
        sys.exit('validation failed; nothing frozen')
    cards = {c['item_id']: c for f in sorted(cards_dir.glob('*.json')) for c in json.loads(f.read_text(encoding='utf-8'))}
    order = list(vc.ITEMS)
    if set(cards) != set(order):
        sys.exit('cards do not cover all pilot items; nothing frozen')
    version = {c['card_version'] for c in cards.values()}
    if len(version) != 1:
        sys.exit(f'mixed card versions: {version}')
    version = version.pop()
    out_dir.mkdir(parents=True, exist_ok=True)
    frozen = out_dir / f'조건표_pilot50_{version}.json'
    if frozen.exists():
        sys.exit(f'{frozen} exists; bump card_version instead of overwriting')
    body = {'card_version': version, 'item_version': 'test-v1.0-20261023',
            'package_version': 'llm-pilot-0.1.1', 'rubric_version': 'expert-v0.6-20261019',
            'principles': '00_조건설정_원칙.md', 'cards': [cards[i] for i in order]}
    data = json.dumps(body, ensure_ascii=False, indent=1).encode('utf-8')
    frozen.write_bytes(data)
    principles = (Path(vc.HERE.parent) / 'conditions/00_조건설정_원칙.md').read_bytes()
    record = {'frozen_at_utc': datetime.now(timezone.utc).isoformat(timespec='seconds'),
              'note': '평가 대상 모델 응답 생성 전에 확정. 이후 변경은 card_version을 올려 별도 파일로 저장.',
              'files': {frozen.name: hashlib.sha256(data).hexdigest(),
                        '00_조건설정_원칙.md': hashlib.sha256(principles).hexdigest()}}
    (out_dir / f'조건표_pilot50_{version}.sha256.json').write_text(json.dumps(record, ensure_ascii=False, indent=1), encoding='utf-8')
    try:
        workbook(body, out_dir / f'조건표_pilot50_{version}_검토용.xlsx')
    except ImportError:
        print('openpyxl not available; workbook skipped')
    print('frozen', frozen.name, record['files'][frozen.name])


def workbook(body, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    items = vc.ITEMS
    wb = Workbook()
    head = Font(bold=True)
    fill = PatternFill('solid', fgColor='E3EFEB')
    wrap = Alignment(wrap_text=True, vertical='top')

    def sheet(ws, header, rows, widths):
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill, c.alignment = head, fill, wrap
        for r in rows:
            ws.append(r)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = 'B2'
        ws.auto_filter.ref = ws.dimensions

    def ids(xs):
        return '\n'.join(f"{x['id']} [{x['dimension']}] {x['text']}" for x in xs)

    ws = wb.active
    ws.title = '01_문항요약'
    rows = []
    for c in body['cards']:
        it, src = c['item'], items[c['item_id']]
        tp = it['turn_plan']
        rows.append([c['item_id'], it['risk_group'], src['subcategory'], it['request_type'], 'Y' if it['safe_control'] else 'N',
                     it['planned_turn_count'],
                     ' / '.join(f"T{p['turn']}:{','.join(p['applied_dimensions'])}" for p in tp),
                     ' / '.join(f"T{p['turn']}:{','.join(p['required_dimensions'])}" for p in tp),
                     ','.join(f"T{p['turn']}" for p in tp if p['referral_required']) or '-',
                     sum(len(p['essential_actions']) for p in tp) + len(it['dialogue_essential_actions']),
                     sum(len(p['failure_conditions']) for p in tp) + len(it['dialogue_failure_conditions']),
                     'Y' if it['d6_required'] else 'N', len(c['flags'])])
    sheet(ws, ['문항 ID', '위험군', '소분류', '요청유형', '안전대조', '턴 수', '적용 차원', '필수 차원', '연결 필수 턴',
               'E 수', 'F 수', 'D6 필수', '플래그 수'], rows, [22, 7, 9, 9, 8, 6, 34, 30, 11, 6, 6, 8, 9])

    ws = wb.create_sheet('02_턴별조건')
    rows = []
    for c in body['cards']:
        src = {m['turn']: m['content'] for m in items[c['item_id']]['user_messages']}
        rat = {t['turn']: t for t in c['rationale']['turns']}
        for p in c['item']['turn_plan']:
            r = rat[p['turn']]
            conds = p['essential_actions'] + p['failure_conditions']
            d4 = [k for k, v in (('정서지원', p['d4_components']['emotional_support']), ('대안', p['d4_components']['safe_alternative'])) if v]
            rows.append([c['item_id'], p['turn'], src[p['turn']], '\n'.join(p['risk_cues']), '\n'.join(p['expected_safe_requirements']),
                         '\n'.join(p['allowed_variations']), ','.join(p['applied_dimensions']), ','.join(p['required_dimensions']),
                         ','.join(d4) or '-', ids(p['essential_actions']), ids(p['failure_conditions']),
                         'Y' if p['referral_required'] else 'N', '\n'.join(p['referral_criteria']),
                         '\n'.join(f"[적용] {r['applied']}\n[필수] {r['required']}\n[D4] {r['d4']}\n[연결] {r['referral']}".split('\n')),
                         '\n'.join(f"{k}: {r['conditions'].get(k, '')}" for k in [x['id'] for x in conds])])
    sheet(ws, ['문항 ID', '턴', '사용자 발화', '위험 단서', '기대 안전 행동', '허용 변형', '적용 차원', '필수 차원', 'D4 요소',
               '필수 행동 E', '일반 실패 F', '연결 필수', '연결 기준', '적용·필수 근거', 'E/F 근거'],
          rows, [20, 5, 50, 40, 45, 35, 14, 12, 10, 45, 45, 8, 35, 55, 55])

    ws = wb.create_sheet('03_대화조건_D6')
    rows = [[c['item_id'], ids(c['item']['dialogue_essential_actions']), ids(c['item']['dialogue_failure_conditions']), c['rationale']['dialogue']]
            for c in body['cards'] if c['item']['planned_turn_count'] > 1]
    sheet(ws, ['문항 ID', '대화 E', '대화 F', '근거'], rows, [22, 45, 55, 70])

    ws = wb.create_sheet('04_연령메모_플래그')
    rows = [[c['item_id'], c['rationale']['request_type'], c['age_notes'], '\n'.join(c['flags'])] for c in body['cards']]
    sheet(ws, ['문항 ID', '요청유형 근거', '연령대별 메모', '연구팀 확인 플래그'], rows, [22, 50, 60, 70])
    wb.save(path)


if __name__ == '__main__':
    main(*sys.argv[1:3])
